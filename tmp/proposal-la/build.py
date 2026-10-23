from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import math

ROOT=Path(__file__).parent
OUT=Path(r'D:/poltekk/SEMESTER 5/PROPOSAL LAPORAN AKHIR/Proposal LA Web Mobile')
OUT.mkdir(exist_ok=True)
TITLE='Sistem Informasi Pengelolaan Donasi Berbasis Web dan Mobile pada Panti Asuhan NU An-Nuur 2 Kota Kediri'
FONT=r'C:/Windows/Fonts/times.ttf'
BOLD=r'C:/Windows/Fonts/timesbd.ttf'

class Diagram:
 def __init__(self,w=1500,h=1200):
  self.im=Image.new('RGB',(w,h),'white');self.d=ImageDraw.Draw(self.im)
 def text(self,xy,t,size=32,bold=False,anchor='mm'):
  self.d.multiline_text(xy,t,font=ImageFont.truetype(BOLD if bold else FONT,size),fill='black',anchor=anchor,align='center',spacing=7)
 def box(self,b,t,size=31,ellipse=False):
  (self.d.ellipse if ellipse else self.d.rounded_rectangle)(b,outline='black',width=3,**({} if ellipse else {'radius':14}))
  self.text(((b[0]+b[2])/2,(b[1]+b[3])/2),t,size)
 def arrow(self,pts,label=None):
  self.d.line(pts,fill='black',width=3)
  a,b=pts[-2:];ang=math.atan2(b[1]-a[1],b[0]-a[0]);l=16
  self.d.polygon([b,(b[0]-l*math.cos(ang-.45),b[1]-l*math.sin(ang-.45)),(b[0]-l*math.cos(ang+.45),b[1]-l*math.sin(ang+.45))],fill='black')
  if label:
   horizontal=abs(b[0]-a[0])>abs(b[1]-a[1])
   pos=((a[0]+b[0])/2, (a[1]+b[1])/2-27) if horizontal else ((a[0]+b[0])/2+52,(a[1]+b[1])/2)
   self.text(pos,label,26)
 def actor(self,x,y,t):
  self.d.ellipse((x-25,y-80,x+25,y-30),outline='black',width=3);self.d.line([(x,y-30),(x,y+40)],fill='black',width=3);self.d.line([(x-45,y),(x+45,y)],fill='black',width=3);self.d.line([(x-40,y+90),(x,y+40),(x+40,y+90)],fill='black',width=3);self.text((x,y+135),t,31)
 def save(self,n):
  p=ROOT/(n+'.png');self.im.save(p);return p

g=Diagram(1500,1050)
g.box((70,70,650,240),'Web publik dan panel pengurus\nBlade · Tailwind CSS · Alpine.js')
g.box((850,70,1430,240),'Aplikasi Android\nFlutter · akun donatur tetap')
g.box((320,380,1180,570),'Backend Laravel\nValidasi · otorisasi · transaksi · REST API /api/v1',34)
g.box((70,770,470,960),'MySQL\nBasis data bersama')
g.box((550,770,950,960),'Midtrans sandbox\nPembayaran dan webhook')
g.box((1030,770,1430,960),'Fonnte dan email\nNotifikasi dan aktivasi')
g.arrow([(360,240),(360,380)],'HTTP(S)');g.arrow([(1140,240),(1140,380)],'JSON')
for x in [270,750,1230]:g.arrow([(750,570),(750,650),(x,650),(x,770)])
g.save('arsitektur')

for n,actor,items in [('usecase-publik','Pengunjung /\ndonatur biasa',['Melihat informasi panti','Mengajukan donasi uang','Melihat status transaksi','Mengajukan bantuan barang']),('usecase-mobile','Donatur tetap',['Mendaftar dan aktivasi akun','Login dan mengelola profil','Berdonasi melalui mobile','Melihat riwayat pribadi','Memantau transparansi dana']),('usecase-pengurus','Pengurus',['Mengelola konten dan kebutuhan','Memverifikasi akun donatur','Mencatat donasi manual','Mengonfirmasi bantuan barang','Membukukan pengeluaran','Mengelola saldo dan laporan'])]:
 g=Diagram(1500,1100);g.d.rectangle((420,50,1450,1050),outline='black',width=3);g.text((935,90),'Sistem informasi pengelolaan donasi',34,True);g.actor(180,490,actor)
 for i,t in enumerate(items):
  y=190+i*140;g.box((650,y,1320,y+105),t,32,True);g.d.line([(225,490),(650,y+52)],fill='black',width=2)
 g.save(n)

def flow(name,steps,notes):
 g=Diagram(1500,1450);g.d.ellipse((725,20,775,70),fill='black');last=(750,70)
 for i,t in enumerate(steps):
  y=115+i*190;g.arrow([last,(750,y)]);g.box((300,y,1200,y+125),t,33);last=(750,y+125)
 g.arrow([last,(750,1350)]);g.d.ellipse((723,1350,777,1404),outline='black',width=4);g.d.ellipse((734,1361,766,1393),fill='black')
 g.save(name)
flow('aktivitas-donasi',['Donatur memilih kategori dan nominal\nWeb tanpa akun atau mobile terautentikasi','Backend memvalidasi dan menyimpan donasi\nStatus awal Menunggu pembayaran','Backend meminta checkout Midtrans\nDonatur menyelesaikan pembayaran sandbox','Midtrans mengirim webhook\nBackend memverifikasi signature dan nominal','Backend memperbarui status secara idempoten\nCallback tidak sah ditolak','Web atau mobile membaca status dari backend\nRiwayat dan laporan diperbarui'],None)
flow('aktivitas-akun',['Donatur mengajukan pendaftaran melalui mobile\nIdentitas dan kredensial divalidasi','Backend membuat akun berstatus menunggu\nPengurus meninjau pengajuan akun','Pengurus menyetujui dan mengirim tautan email\nAkun belum aktif sebelum tautan ditukarkan','Donatur membuka tautan aktivasi yang sah\nToken kedaluwarsa atau terpakai ditolak','Backend memverifikasi email dan mengaktifkan akun\nDonatur login untuk memperoleh token akses','Donatur memakai layanan pribadi\nLogout atau pencabutan akses menonaktifkan sesi'],None)
flow('aktivitas-pengeluaran',['Pengurus mengisi tanggal, kategori dan nominal\nMengunggah bukti transaksi bila tersedia','Backend menyimpan pengeluaran sebagai Draft\nSaldo belum berkurang','Pengurus memeriksa catatan dan memilih pembukuan\nValidasi gagal mengembalikan pesan koreksi','Backend mengubah status menjadi Dibukukan\nMencatat pelaksana dan waktu pembukuan','Laporan menghitung pengeluaran satu kali\nPenerimaan dan saldo memakai periode yang sama','Mobile menampilkan ringkasan dan rincian terpilih\nBukti hanya dapat dibuka jika diizinkan publik'],None)

g=Diagram(1500,1300)
boxes={'users':(40,50,610,270),'donatur':(890,50,1460,270),'donasi':(890,480,1460,770),'payment_transactions':(890,1000,1460,1220),'kategori_donasi':(40,480,610,700),'donor_account_audits':(40,950,610,1220)}
texts={'users':'users\nPK id · name · email\nrole · account_status','donatur':'donatur\nPK id · FK user_id (unik, opsional)\nnama · whatsapp · jenis','donasi':'donasi\nPK id · FK donatur_id\nFK owner_user_id (opsional)\nFK kategori_donasi_id\nnominal · biaya · status · sumber','payment_transactions':'payment_transactions\nPK id · FK donasi_id (unik)\nReferensi dan status provider','kategori_donasi':'kategori_donasi\nPK id · nama','donor_account_audits':'donor_account_audits\nPK id · FK donatur_id\nFK actor_user_id · action'}
for k,b in boxes.items():g.box(b,texts[k],30)
g.arrow([(610,150),(890,150)],'1 : 0..1');g.arrow([(1175,270),(1175,480)],'1 : N');g.arrow([(610,590),(890,590)],'1 : N');g.arrow([(1175,770),(1175,1000)],'1 : 0..1');g.arrow([(330,270),(330,380),(800,380),(800,690),(890,690)],'pemilik 1 : N');g.arrow([(890,210),(750,210),(750,1080),(610,1080)],'1 : N')
g.save('erd-donasi')
g=Diagram(1500,1350)
for b,t in [((40,70,600,280),'donatur\nPK id'),((900,70,1460,280),'bantuan\nPK id · FK donatur_id · status'),((900,480,1460,720),'bantuan_items\nPK id · FK bantuan_id\nFK kategori_bantuan_id\nnama_barang · jumlah · satuan'),((40,480,600,720),'kategori_bantuan\nPK id · nama'),((40,950,600,1190),'kategori_pengeluaran\nPK id · nama'),((900,950,1460,1190),'pengeluaran\nPK id · FK kategori_pengeluaran_id\nFK created_by / posted_by (users)\nnominal · status · is_public')]:g.box(b,t,30)
g.arrow([(600,170),(900,170)],'1 : N');g.arrow([(1180,280),(1180,480)],'1 : N');g.arrow([(600,600),(900,600)],'1 : N');g.arrow([(600,1070),(900,1070)],'1 : N');g.text((750,1290),'Saldo pembukaan terkait users melalui disahkan_oleh (N : 1)',31)
g.save('erd-bantuan')

g=Diagram(1500,1050);g.box((30,30,1470,1020),'',30);g.d.rectangle((30,30,1470,135),outline='black',width=3);g.text((750,80),'AN-NUUR 2     Profil     Kegiatan     Kebutuhan     Donasi',34,True)
g.box((90,200,1410,445),'Profil panti dan informasi kegiatan\n[ Lihat kebutuhan ]       [ Salurkan donasi ]',35)
for x,t in [(90,'Kebutuhan\nTarget dan realisasi'),(550,'Kegiatan\nFoto dan ringkasan'),(1010,'Bantuan barang\nForm pengajuan')]:g.box((x,520,x+400,760),t,32)
g.text((750,900),'Kontak dan lokasi panti · Informasi layanan',32);g.save('ui-web')
g=Diagram(1500,1050);g.box((30,30,1470,1020),'');g.box((50,50,345,1000),'Panel pengurus\n\nDashboard\nDonatur\nDonasi\nBantuan\nPengeluaran\nArus kas\nLaporan',32)
g.text((890,100),'PENGELOLAAN DONASI DAN ARUS KAS',34,True)
for x,t in [(390,'Penerimaan\nbersih'),(750,'Pengeluaran\ndibukukan'),(1110,'Saldo\ntercatat')]:g.box((x,180,x+310,340),t)
g.box((400,410,1420,810),'Filter periode dan status\n\nTanggal     Keterangan     Nominal     Status\n\nDaftar transaksi / pengeluaran\n\n[ Detail ]  [ Bukukan ]  [ PDF ]  [ Excel ]',33);g.text((900,920),'Peringatan biaya belum tersedia / saldo sementara',30);g.save('ui-admin')
g=Diagram(1500,1200)
for x,title,body in [(40,'Beranda','Total donasi pribadi\n\n[ Donasi sekarang ]\n\nRingkasan dana panti\nPenerimaan bersih\nPengeluaran\nSaldo'),(540,'Donasi','Kategori donasi\n[ Pilih kategori ]\n\nNominal\n[ Pilihan / lainnya ]\n\n[ Lanjut bayar ]\n\nStatus dari server'),(1040,'Transparansi','Bulan dan tahun\n[ Pilih periode ]\n\nRingkasan arus kas\n\nDaftar pengeluaran\n[ Detail ]\n[ Bukti publik ]')]:
 g.box((x,40,x+420,1120),'');g.text((x+210,125),title,38,True);g.d.line([(x+20,180),(x+400,180)],fill='black',width=2);g.text((x+210,600),body,32);g.text((x+210,1035),'Beranda · Riwayat\nDana · Profil',27)
g.save('ui-mobile')

doc=Document()
for s in doc.sections:
 s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(3);s.bottom_margin=Cm(3);s.left_margin=Cm(4);s.right_margin=Cm(3);s.header_distance=Cm(1.5);s.footer_distance=Cm(1.5)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption','TOC 1','TOC 2','TOC 3']:
 try:st=doc.styles[name]
 except KeyError:continue
 st.font.name='Times New Roman';st.font.size=Pt(12);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(0);st.paragraph_format.line_spacing=1.5
doc.styles['Normal'].paragraph_format.first_line_indent=Cm(1)
doc.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for n in ['Heading 1','Heading 2','Heading 3']:
 st=doc.styles[n];st.font.bold=True;st.paragraph_format.first_line_indent=Cm(0);st.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;doc.styles['Heading 1'].paragraph_format.space_after=Pt(18)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(12)
doc.styles['Heading 3'].paragraph_format.space_before=Pt(6)

def p(t='',style=None):return doc.add_paragraph(t,style)
def center(t,size=12,bold=False,after=12):
 q=p();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.first_line_indent=Cm(0);q.paragraph_format.space_after=Pt(after);r=q.add_run(t);r.font.size=Pt(size);r.bold=bold;return q
def h(t,lv=2):return p(t,f'Heading {lv}')
def paras(t):
 for x in t.strip().split('\n\n'):p(x.replace('\n',' '))
def item(t):
 q=p(t);q.paragraph_format.first_line_indent=Cm(-.5);q.paragraph_format.left_indent=Cm(.5);q.paragraph_format.space_after=Pt(4)
def page():doc.add_page_break()
def field(q,code):
 a=OxmlElement('w:fldSimple');a.set(qn('w:instr'),code);q._p.append(a)
def numbering(sec,fmt='decimal',start=None,chapter=False):
 for part in [sec.header,sec.footer,sec.first_page_header,sec.first_page_footer]:
  part.is_linked_to_previous=False
  for q in part.paragraphs:q.clear()
 for old in list(sec._sectPr.findall(qn('w:pgNumType'))):sec._sectPr.remove(old)
 e=OxmlElement('w:pgNumType');e.set(qn('w:fmt'),fmt)
 if start:e.set(qn('w:start'),str(start))
 sec._sectPr.append(e);sec.different_first_page_header_footer=chapter
 q=sec.header.paragraphs[0] if chapter else sec.footer.paragraphs[0];q.alignment=WD_ALIGN_PARAGRAPH.RIGHT if chapter else WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.first_line_indent=Cm(0);field(q,'PAGE')
 if chapter:
  q=sec.first_page_footer.paragraphs[0];q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.first_line_indent=Cm(0);field(q,'PAGE')
def chapter(t,start=None):
 sec=doc.add_section(WD_SECTION_START.NEW_PAGE);numbering(sec,start=start,chapter=True);h(t,1)
def table(title,headers,rows,widths=None):
 q=center(title,11,False,3);q.paragraph_format.keep_with_next=True
 t=doc.add_table(rows=1, cols=len(headers));t.style='Table Grid';t.autofit=False
 for i,s in enumerate(headers):t.rows[0].cells[i].text=s
 for row in rows:
  c=t.add_row().cells
  for i,s in enumerate(row):c[i].text=str(s)
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trPr.append(ns)
  if ri==0:trPr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   if widths:c.width=Cm(widths[ci])
   for q in c.paragraphs:
    q.paragraph_format.first_line_indent=Cm(0);q.paragraph_format.line_spacing=1;q.paragraph_format.space_after=Pt(4);q.paragraph_format.space_before=Pt(4);q.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for r in q.runs:r.font.size=Pt(11);r.bold=ri==0
 q=p('Sumber: Rancangan penulis');q.paragraph_format.first_line_indent=Cm(0);q.paragraph_format.line_spacing=1;q.paragraph_format.space_after=Pt(10)
 for r in q.runs:r.font.size=Pt(10)
 return t
def fig(name,caption,width=14):
 q=p();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.first_line_indent=Cm(0);q.paragraph_format.keep_with_next=True;q.add_run().add_picture(str(ROOT/(name+'.png')),width=Cm(width))
 q=center(caption,11,False,3);q.paragraph_format.line_spacing=1
 q=center('Sumber: Rancangan penulis',10,False,12);q.paragraph_format.line_spacing=1

def cover():
 center('PROPOSAL LAPORAN AKHIR',14,True,28)
 titlep=center(TITLE.upper(),14,True,30);titlep.style=doc.styles['Title'];titlep.paragraph_format.line_spacing=1.5
 center('Oleh:',12,False,0);center('NOOR ACHMAD ROZAQ RAMDHANI\nNIM. 243107030100',12,True,30)
 q=center('',12,False,30);q.add_run().add_picture(str(ROOT/'guide-33-0.jpg'),width=Cm(4),height=Cm(4))
 center('PROGRAM STUDI D3 MANAJEMEN INFORMATIKA\nJURUSAN TEKNOLOGI INFORMASI\nPSDKU POLINEMA DI KOTA KEDIRI\nPOLITEKNIK NEGERI MALANG\n2026',14,True,0)
cover()
s=doc.add_section(WD_SECTION_START.NEW_PAGE);numbering(s,'lowerRoman',1);cover()
page();center('HALAMAN PENGESAHAN',14,True,18);center(TITLE.upper(),12,True,18)
center('Oleh:\nNoor Achmad Rozaq Ramdhani\nNIM. 243107030100',12,False,18)
center('Proposal Laporan Akhir ini telah diuji pada tanggal\n........................................................',12,False,18)
center('Disetujui oleh:',12,False,12)
for label,name,nip in [('Pembimbing','Fery Sofian Efendi, S.Kom., M.Cs.','197809202024211003'),('Penguji I','Fikha Rizky Aulia, M.Pd.','199306202025212061'),('Penguji II','Kunti Eliyen, S.Kom., M.T.','198905282024212028')]:
 q=p(f'{label}\n\n{name}\nNIP. {nip}');q.alignment=WD_ALIGN_PARAGRAPH.LEFT;q.paragraph_format.first_line_indent=Cm(0);q.paragraph_format.line_spacing=1;q.paragraph_format.space_after=Pt(14)
center('Mengetahui,\nKoordinator Program Studi D3 Manajemen Informatika\n\n\n........................................................\nNIP/NIDN. ........................................',12,False,0)
page();center('DAFTAR ISI',14,True,18);field(p(),'TOC \\o "1-3" \\h \\z \\u')

chapter('BAB I\nPENDAHULUAN',1)
h('1.1 Latar Belakang')
paras('''Pengelolaan donasi pada panti asuhan mencakup pencatatan pemberi donasi, penerimaan uang dan barang, penggunaan dana, serta penyampaian laporan. Keterkaitan proses tersebut membutuhkan informasi yang dapat ditelusuri agar pengurus dapat mengetahui transaksi yang sudah diterima dan dana yang telah digunakan. Penelitian pada sistem administrasi panti menunjukkan bahwa pencatatan transaksi dan penyusunan laporan merupakan kebutuhan penting dalam pengelolaan lembaga sosial (Salim, 2020).

Panti Asuhan NU An-Nuur 2 Kota Kediri menjadi objek pengembangan sistem informasi dalam penelitian ini. Kebutuhan awal yang dirumuskan mencakup penyajian informasi panti, pengelolaan donatur, donasi uang, bantuan barang, pengeluaran, dan laporan. Informasi tersebut perlu dihubungkan agar pencatatan oleh pengurus dan informasi yang diterima donatur merujuk pada sumber data yang sama. Alur administrasi yang saat ini digunakan, bentuk dokumen, serta kendala operasional akan dikonfirmasi melalui observasi dan wawancara dengan pengurus.

Layanan donasi juga perlu mempertimbangkan perbedaan kebutuhan pengguna. Donatur biasa memerlukan akses melalui web tanpa kewajiban membuat akun. Sementara itu, donatur tetap membutuhkan layanan yang memungkinkan pemberian donasi dan penelusuran riwayat pribadi secara berulang melalui aplikasi mobile. Pembagian tersebut menuntut pengelolaan identitas dan hak akses yang jelas. Kesamaan nomor kontak tidak dapat langsung dijadikan dasar untuk memberikan akses terhadap seluruh riwayat transaksi.

Penelitian terdahulu telah menghasilkan sistem web yang mendukung publikasi panti dan penggalangan donasi. Sistem e-Panti menyajikan informasi profil serta layanan donasi, sedangkan sistem penggalangan dana di Garut memberikan peluang pengembangan menuju aplikasi Android (Mauliana, Firmansyah, & Sutardi, 2019; Romansyah, Mulyani, & Kurniadi, 2019). Penelitian lain membahas pencatatan penerimaan, pengeluaran, dan laporan keuangan panti sebagai satu rangkaian administrasi (Yunita et al., 2022; Destriani & Bahri, 2026).

Transparansi penggunaan donasi menjadi pertimbangan tambahan dalam rancangan sistem. Sistem berbasis Laravel pada Panti Asuhan Mawaddah Warohmah menghubungkan pengelolaan administrasi dengan informasi penggunaan dana, sedangkan Panti Peduli mengusulkan pengembangan aplikasi mobile sebagai salah satu arah pengembangan lanjutan (Sintiya et al., 2025; Safira, Suryantari, & Ramadhan, 2024). Temuan tersebut menjadi dasar untuk merancang akses informasi dana yang terhubung dengan proses pencatatan pengurus, tanpa menyatakan bahwa seluruh penelitian terdahulu belum memiliki layanan yang sama.

Penelitian ini mengusulkan Sistem Informasi Pengelolaan Donasi Berbasis Web dan Mobile pada Panti Asuhan NU An-Nuur 2 Kota Kediri. Web digunakan untuk publikasi, donasi tanpa akun, bantuan barang, dan pengelolaan administrasi pengurus. Aplikasi Android digunakan oleh donatur tetap untuk berdonasi, melihat riwayat pribadi, serta memantau ringkasan dana dan pengeluaran. Backend Laravel dan basis data MySQL menjadi pusat pengolahan data, sedangkan aplikasi Flutter berkomunikasi melalui API. Integrasi pembayaran diuji menggunakan Midtrans sandbox.

Sistem diharapkan membantu pengurus menelusuri penerimaan dan penggunaan dana serta memudahkan donatur mengakses informasi sesuai haknya. Ketercapaian tujuan akan diperiksa melalui pengujian fungsi, konsistensi data web dan mobile, serta pengujian penerimaan pengguna. Dampak terhadap efisiensi dan kemudahan penggunaan akan dievaluasi berdasarkan bukti pengujian, bukan diasumsikan hanya karena sistem telah dibangun.''')
h('1.2 Rumusan Masalah')
for x in ['1. Bagaimana mengembangkan sistem informasi berbasis web dan mobile yang mengintegrasikan pengelolaan donatur, donasi uang, bantuan barang, pengeluaran, dan laporan pada Panti Asuhan NU An-Nuur 2 Kota Kediri?','2. Bagaimana menyediakan layanan donasi tanpa akun melalui web dan layanan donatur tetap melalui mobile dengan pengendalian hak akses serta riwayat transaksi pribadi?','3. Bagaimana menyajikan informasi penerimaan, pengeluaran, dan saldo tercatat yang konsisten antara web dan mobile serta menguji kesesuaian sistem dengan kebutuhan pengguna?']:item(x)
h('1.3 Tujuan')
for x in ['1. Menghasilkan sistem informasi web dan mobile yang mengintegrasikan data donatur, donasi uang, bantuan barang, pengeluaran, dan laporan.','2. Menyediakan donasi web tanpa akun dan layanan mobile donatur tetap yang meliputi aktivasi akun, donasi, riwayat pribadi, serta pengelolaan profil sesuai hak akses.','3. Menyediakan informasi arus dana yang konsisten antarplatform dan mengevaluasi fungsi, integrasi, serta penerimaan sistem oleh pengurus dan donatur.']:item(x)
h('1.4 Batasan Masalah')
for x in ['1. Objek penelitian dibatasi pada Panti Asuhan NU An-Nuur 2 Kota Kediri; sistem tidak menjadi platform penggalangan dana bagi banyak panti.','2. Web mencakup profil, kegiatan, berita, galeri, kebutuhan, donasi uang tanpa akun, pengajuan bantuan barang, dan panel pengurus.','3. Mobile ditujukan untuk Android menggunakan Flutter. Layanan pribadi diperuntukkan bagi donatur tetap dengan akun yang disetujui pengurus dan diaktifkan melalui email.','4. Pengurus memiliki tingkat kewenangan aplikasi yang sama. Pimpinan atau bendahara tidak dijadikan peran akun terpisah dalam penelitian ini.','5. Donasi uang mencakup pembayaran online dan pencatatan penerimaan tunai atau transfer langsung oleh pengurus. Integrasi pembayaran dibatasi pada lingkungan Midtrans sandbox.','6. Bantuan barang diajukan melalui web dan dicatat diterima secara keseluruhan oleh pengurus. Penerimaan sebagian, manajemen stok lanjutan, dan bantuan barang melalui mobile tidak termasuk.','7. Informasi keuangan meliputi saldo pembukaan, penerimaan bruto, biaya tercatat, penerimaan bersih, pengeluaran dibukukan, serta saldo akhir tercatat. Sistem bukan aplikasi akuntansi lengkap atau rekonsiliasi bank otomatis.','8. Pengeluaran memiliki status Draft, Dibukukan, dan Dibatalkan. Bukti privat dibatasi aksesnya; hanya bukti yang diberi izin publik yang tersedia melalui layanan transparansi.','9. Transparansi menyajikan agregat dana dan pengeluaran, bukan penelusuran setiap rupiah dari seorang donatur ke pengeluaran tertentu.','10. Pembayaran berulang otomatis, pengembalian dana otomatis, push notification, iOS, publikasi toko aplikasi, serta fitur kunjungan dan buku tamu tidak termasuk. Notifikasi WhatsApp digunakan sebagai layanan pendukung sesuai konfigurasi.','11. Pengujian mencakup fungsi, hak akses, integrasi web–API–mobile, konsistensi laporan, serta penerimaan pengguna. Data demonstrasi dipisahkan dari laporan operasional panti.']:item(x)
h('1.5 Sistematika Penulisan')
paras('''Proposal disusun dalam empat bab. Bab I Pendahuluan menjelaskan latar belakang, rumusan masalah, tujuan, batasan masalah, dan sistematika penulisan. Bab II Tinjauan Pustaka menguraikan penelitian yang relevan dan konsep yang mendasari pengembangan sistem. Bab III Analisis dan Perancangan membahas analisis masalah, pengumpulan data, metode pengembangan, kebutuhan sistem, diagram proses, rancangan data, arsitektur, antarmuka, dan rencana pengujian. Bab IV Jadwal Kegiatan memuat jadwal penelitian serta pembagian tanggung jawab pelaksanaan. Bagian akhir memuat daftar pustaka yang dirujuk dalam proposal.''')

chapter('BAB II\nTINJAUAN PUSTAKA')
h('2.1 Penelitian yang Relevan')
paras('''Sepuluh artikel digunakan untuk menelaah publikasi informasi panti, pengelolaan donasi, administrasi keuangan, dan evaluasi pengguna. Kajian diarahkan pada kemampuan sistem yang dilaporkan peneliti serta relevansinya terhadap kebutuhan web dan mobile pada objek penelitian ini.''')
studies=[
('2.1.1 Sistem e-Panti','Sistem e-Panti dikembangkan menggunakan Waterfall, pemodelan UML, PHP, dan MySQL. Web menyajikan profil, informasi anak asuh, kegiatan, serta cara berdonasi. Evaluasi menggunakan Black Box dan USE Questionnaire. Artikel melaporkan nilai 85,50% pada pertanyaan tertentu dalam uji daya guna, sehingga angka tersebut tidak digunakan sebagai ukuran keberhasilan sistem pada penelitian ini. Kajian ini menjadi acuan penggabungan informasi publik dengan layanan donasi (Mauliana, Firmansyah, & Sutardi, 2019).'),
('2.1.2 Penggalangan Dana Sosial Panti Asuhan','Penelitian di Kota Garut menghasilkan sistem penggalangan dana berbasis web dengan Waterfall dan pemodelan terstruktur. Fungsi diuji menggunakan Black Box. Saran pengembangan menyebut aplikasi Android dan penyempurnaan tampilan. Relevansinya adalah penyediaan kanal donasi online, sedangkan penelitian ini menambahkan layanan aplikasi Android yang terhubung dengan pengelolaan web pada satu panti (Romansyah, Mulyani, & Kurniadi, 2019).'),
('2.1.3 Administrasi Panti Asuhan','Sistem administrasi dirancang mengikuti proses bisnis umum panti, mencakup data, penerimaan sumbangan, pengeluaran, dan laporan. Penelitian mengusulkan pengembangan lanjutan berupa riwayat logistik dan laporan sesuai standar akuntansi. Kajian ini mendukung perlunya keterkaitan data transaksi dan laporan. Ruang lingkup penelitian sekarang dipersempit pada pengelolaan donasi, sehingga administrasi lengkap anak asuh tidak diadopsi (Salim, 2020).'),
('2.1.4 Sistem Geografis dan Donasi Online','Sistem untuk panti asuhan di Medan menggabungkan peta lokasi dengan layanan donasi online dan informasi penyaluran. Kualitas sistem dievaluasi menggunakan WebQual 4.0 yang mencakup kegunaan, kualitas informasi, interaksi layanan, dan kualitas visual. Kajian ini menekankan bahwa penyajian informasi perlu dinilai dari sisi pengguna. Penelitian sekarang berfokus pada pengelolaan satu panti dan integrasi data web–mobile, bukan pemetaan banyak panti (Chandra, Hendrik, & Leonardo, 2022).'),
('2.1.5 Pencatatan Keuangan Harapan Karomah','Sistem administrasi keuangan pada Yayasan Panti Asuhan Harapan Karomah dikembangkan melalui Waterfall dengan PHP, MySQL, dan CodeIgniter. Sistem mengolah data keuangan serta menghasilkan laporan; pengujian yang dilaporkan menunjukkan fungsi berjalan sesuai rancangan. Penelitian ini menjadi acuan hubungan pencatatan penerimaan, pengeluaran, dan pelaporan. Pada sistem yang diusulkan, informasi tersebut juga disajikan melalui mobile dengan pembatasan bukti transaksi (Yunita et al., 2022).'),
('2.1.6 Manajemen Donasi dan Penerimaan Teknologi','Penelitian manajemen donasi pada panti asuhan di Tangerang menggunakan Technology Acceptance Model untuk menelaah penerimaan pengguna. Data dikumpulkan melalui observasi, wawancara, kuesioner, dan studi pustaka. Kajian ini menunjukkan pentingnya menilai kemudahan serta manfaat yang dirasakan pengguna. Penelitian sekarang menggunakan pengujian penerimaan berbasis tugas dan tanggapan pengguna, tanpa mengklaim menerapkan pengujian hipotesis TAM yang sama (Kusnadi & Daniawan, 2024).'),
('2.1.7 Platform Panti Peduli','Panti Peduli menyediakan pencarian panti, registrasi, donasi online, dan pelaporan melalui web. Pengembangan menggunakan Waterfall dengan pengumpulan data kualitatif. Artikel mengusulkan pengembangan mobile serta peningkatan keamanan dan transparansi. Penelitian ini mengadaptasi gagasan akses donasi dan informasi, tetapi membatasi objek pada satu panti dan menghubungkan aplikasi donatur tetap dengan administrasi pengurus (Safira, Suryantari, & Ramadhan, 2024).'),
('2.1.8 Transparansi Donasi dengan Extreme Programming','Sistem pada Panti Asuhan Mawaddah Warohmah dikembangkan menggunakan Extreme Programming dengan tahap perencanaan, desain, pengkodean, dan pengujian. Laravel digunakan untuk layanan donatur dan administrasi pengurus. Artikel melaporkan pengujian Black Box sesuai fungsi serta rata-rata kepuasan UAT 96,8%. Relevansinya terletak pada pengelolaan dana dan evaluasi pengguna; hasil UAT tersebut merupakan hasil penelitian terdahulu, bukan target atau hasil penelitian ini (Sintiya et al., 2025).'),
('2.1.9 Pengelolaan Donasi Al-Rasyid','Sistem pada Panti Asuhan Al-Rasyid menggunakan Waterfall untuk mengelola data donatur, penerimaan donasi, pengeluaran, serta laporan. Artikel menjelaskan otomatisasi pencatatan dan kemudahan donasi online. Kajian ini menjadi acuan bahwa donasi dan pengeluaran perlu diolah dalam rangkaian yang sama. Penelitian sekarang memperinci pemisahan bruto dan bersih, status pembukuan, kepemilikan riwayat, serta penyajian informasi lintas platform (Destriani & Bahri, 2026).'),
('2.1.10 Analisis Pendataan Anak Asuh','Kajian di Panti Asuhan Muhammadiyah Slawi menggunakan pendekatan deskriptif kualitatif dan SWOT untuk menganalisis pendataan manual serta semi-digital. Hasilnya menekankan kebutuhan validasi, kontrol akses, pencadangan, dan kesiapan staf. Artikel merupakan kajian analisis pendataan, sehingga tidak diperlakukan sebagai bukti implementasi aplikasi donasi. Relevansinya bagi penelitian ini adalah pendekatan identifikasi kebutuhan dan pengendalian kualitas data (Ujianti & Al-Farisy, 2026).')]
for title,body in studies:h(title,3);p(body)
h('2.1.11 Posisi Penelitian yang Diusulkan',3)
paras('''Berdasarkan kajian tersebut, penelitian diarahkan pada integrasi proses donasi dan penggunaan dana pada Panti Asuhan NU An-Nuur 2 Kota Kediri. Kontribusi yang diusulkan adalah pemisahan layanan web tanpa akun dan mobile donatur tetap, pengendalian kepemilikan riwayat, serta penyajian penerimaan bersih dan pengeluaran dibukukan dari sumber data bersama. Kebaruan dinyatakan pada rancangan dan penerapannya untuk kebutuhan objek penelitian, bukan sebagai klaim bahwa fitur serupa belum pernah dibuat.

Kesesuaian informasi menjadi perhatian utama. Status pembayaran harus berasal dari backend yang memverifikasi pemberitahuan penyedia pembayaran, sementara nominal laporan harus menggunakan aturan dan periode yang sama di web dan mobile. Informasi pengeluaran dapat mendukung transparansi apabila data yang dipublikasikan tetap mempertimbangkan kerahasiaan bukti dan identitas donatur.''')
h('2.2 Landasan Teori')
theories=[('2.2.1 Sistem Informasi Pengelolaan Donasi','Sistem informasi pengelolaan donasi menghubungkan pencatatan donatur, penerimaan, pengeluaran, dan penyusunan informasi untuk pengguna. Dalam konteks panti, transaksi yang tercatat perlu dapat diolah menjadi laporan sesuai kebutuhan pengurus. Keterpaduan tersebut menjadi dasar pemisahan data master, transaksi, dan laporan dalam penelitian ini (Salim, 2020; Destriani & Bahri, 2026).'),('2.2.2 Donasi Uang dan Bantuan Barang','Donasi uang dicatat sebagai nominal penerimaan, sedangkan bantuan barang memerlukan nama barang, jumlah, satuan, dan status penerimaan. Perbedaan karakter data menyebabkan kedua alur dipisahkan. Penerimaan barang tidak langsung dijumlahkan ke saldo uang. Kebutuhan pencatatan sumbangan uang dan barang juga terdapat pada kajian administrasi panti (Salim, 2020).'),('2.2.3 Transparansi dan Pelaporan Dana','Transparansi dalam rancangan ini diwujudkan melalui informasi penerimaan, penggunaan, dan saldo tercatat yang dapat ditelusuri ke transaksi pendukung. Pengeluaran yang belum dibukukan tidak menjadi pengurang saldo. Kajian keuangan dan transparansi panti mendukung keterkaitan pencatatan transaksi dengan penyajian laporan (Yunita et al., 2022; Sintiya et al., 2025).'),('2.2.4 Aplikasi Web dan Mobile','Web menjadi kanal akses melalui peramban untuk masyarakat dan pengurus. Aplikasi mobile menjadi kanal terpasang pada Android bagi donatur tetap. Kedua kanal menggunakan backend yang sama agar perubahan transaksi dapat dibaca secara konsisten. Rancangan ini menindaklanjuti arah pengembangan Android dan mobile yang disebut dalam penelitian penggalangan dana serta Panti Peduli (Romansyah, Mulyani, & Kurniadi, 2019; Safira, Suryantari, & Ramadhan, 2024).'),('2.2.5 Teknologi Implementasi','Laravel digunakan untuk menangani proses aplikasi, validasi, autentikasi, dan penyediaan API. Blade, Tailwind CSS, dan Alpine.js digunakan pada antarmuka web. Flutter digunakan untuk antarmuka Android, sedangkan MySQL menyimpan data bersama. Pilihan teknologi ini merupakan keputusan implementasi penelitian. Penggunaan Laravel dalam pengelolaan panti juga ditemukan pada penelitian transparansi donasi, tetapi artikel tersebut tidak dijadikan sumber spesifikasi Flutter atau API (Sintiya et al., 2025).'),('2.2.6 API dan Pengendalian Akses','Application Programming Interface atau API menjadi antarmuka pertukaran permintaan dan data antara aplikasi mobile dan backend. Rancangan menggunakan JSON dan token autentikasi untuk layanan pribadi. Backend menentukan pemilik transaksi dari akun yang terautentikasi, bukan dari identitas yang bebas dikirim klien. Kontrol akses dan validasi data sejalan dengan kebutuhan kualitas pengelolaan informasi yang dibahas pada analisis pendataan panti (Ujianti & Al-Farisy, 2026).'),('2.2.7 Pemodelan Sistem','Use case digunakan untuk menunjukkan hubungan aktor dengan fungsi, activity diagram menggambarkan urutan kegiatan, dan diagram relasi menunjukkan hubungan data. Pemodelan membantu menghubungkan kebutuhan pengguna dengan rancangan implementasi. Penggunaan UML pada sistem e-Panti dan pemodelan pada sistem donasi menjadi acuan penyusunan rancangan yang dapat ditinjau sebelum implementasi (Mauliana, Firmansyah, & Sutardi, 2019; Chandra, Hendrik, & Leonardo, 2022).'),('2.2.8 Metode Waterfall','Waterfall mengorganisasikan pekerjaan menjadi analisis kebutuhan, desain, implementasi, pengujian, dan pemeliharaan. Setiap tahap menghasilkan dokumen atau perangkat lunak yang menjadi dasar tahap berikutnya. Model ini dipilih untuk menjaga keterlacakan kebutuhan dan rancangan proposal. Perubahan kebutuhan tetap dicatat dan dinilai dampaknya sebelum rancangan serta implementasi diperbarui (Romansyah, Mulyani, & Kurniadi, 2019; Yunita et al., 2022).'),('2.2.9 Pengujian Fungsi dan Penerimaan Pengguna','Pengujian Black Box memeriksa keluaran sistem berdasarkan masukan dan kondisi tertentu tanpa bergantung pada pemeriksaan struktur kode. Pengujian penerimaan melibatkan pengguna untuk menilai keberhasilan menjalankan tugas. Penggunaan keduanya pada sistem panti menjadi dasar pemisahan pemeriksaan fungsi dengan penilaian pengalaman pengguna dalam penelitian ini (Mauliana, Firmansyah, & Sutardi, 2019; Sintiya et al., 2025).')]
for title,body in theories:h(title,3);p(body)

chapter('BAB III\nANALISIS DAN PERANCANGAN')
h('3.1 Analisis Permasalahan')
paras('''Analisis dilakukan pada rangkaian layanan donasi, mulai dari penyampaian informasi panti, penerimaan uang atau barang, pencatatan pengeluaran, sampai penyusunan laporan. Berdasarkan kebutuhan awal, pengurus membutuhkan data terpusat, sedangkan donatur membutuhkan akses informasi dan status transaksi. Proses berjalan akan dipetakan melalui observasi dan wawancara agar rancangan sesuai pembagian pekerjaan yang berlaku di panti.

Penelusuran lapangan akan memeriksa siapa yang menerima donasi, bagaimana bukti pembayaran diverifikasi, kapan barang dinyatakan diterima, serta bagaimana laporan dana dibuat. Dokumen yang diperiksa meliputi catatan penerimaan, daftar donatur, bukti pengeluaran, dan contoh laporan yang diizinkan pengurus. Hasilnya digunakan untuk memvalidasi kebutuhan, bukan untuk mengasumsikan seluruh proses panti masih manual.''')
table('Tabel 3.1 Pemetaan kebutuhan dan penyelesaian',['Kebutuhan awal','Risiko yang perlu diperiksa','Penyelesaian'],[
('Informasi panti dan kebutuhan','Informasi tersebar atau tidak mutakhir','Web publik dengan pengelolaan konten oleh pengurus.'),('Penerimaan uang dan barang','Status pembayaran atau penerimaan tidak tertelusuri','Pemisahan transaksi donasi dan bantuan beserta riwayat status.'),('Riwayat donatur tetap','Transaksi salah dikaitkan ke akun','Kepemilikan transaksi berdasarkan akun atau verifikasi pengurus.'),('Pengeluaran dan laporan','Angka laporan berbeda antarpenyajian','Satu aturan perhitungan pada backend dan filter periode yang sama.'),('Transparansi dana','Bukti sensitif terpublikasi','Pembatasan bukti privat dan penyajian informasi publik terpilih.')],[4,4.5,5.5])
h('3.2 Analisis Pemecahan Masalah')
h('3.2.1 Metode Pengambilan Data',3)
paras('''Pengumpulan data direncanakan melalui observasi, wawancara, dokumentasi, dan studi pustaka. Informan utama adalah pengurus yang menangani administrasi dan penerimaan atau penggunaan donasi. Donatur dilibatkan untuk mengidentifikasi kebutuhan informasi, pengalaman berdonasi, serta kemampuan menggunakan aplikasi Android.

Observasi digunakan untuk mencatat alur penerimaan uang, barang, dan pengeluaran. Wawancara menggali pembagian tanggung jawab, aturan status transaksi, klasifikasi donatur, persetujuan akun, serta kebijakan publikasi bukti. Dokumentasi digunakan untuk memeriksa struktur data dan bentuk laporan. Studi pustaka dilakukan terhadap sepuluh artikel yang dibahas pada Bab II.

Setiap hasil pengumpulan data akan dirangkum dalam daftar kebutuhan yang diperiksa kembali bersama pengurus. Perbedaan antara keterangan wawancara dan dokumen ditindaklanjuti sebelum kebutuhan ditetapkan. Pengambilan salinan dokumen dilakukan dengan izin; identitas dan informasi sensitif disamarkan pada bahan pengujian dan laporan akademik.''')
table('Tabel 3.2 Rencana pertanyaan dan bukti kebutuhan',['Pokok pertanyaan','Bukti yang dicari'],[('Bagaimana donasi diterima dan disahkan?','Alur penerimaan, contoh bukti, dan pihak yang berwenang.'),('Bagaimana pengeluaran dicatat dan dikoreksi?','Catatan pengeluaran, bukti, status pembukuan, dan alur koreksi.'),('Siapa yang memperoleh akun donatur tetap?','Kriteria klasifikasi, persetujuan, aktivasi, serta pencabutan akses.'),('Informasi apa yang boleh dilihat donatur?','Batas publikasi rincian, identitas, dan bukti pengeluaran.'),('Bagaimana laporan dipakai dan diperiksa?','Periode, komponen saldo, format keluaran, dan pemeriksa laporan.')],[7,7])
h('3.2.2 Metode Pengembangan Sistem',3)
paras('''Pengembangan menggunakan Waterfall dengan keluaran yang dapat ditinjau pada setiap tahap. Analisis menghasilkan kebutuhan serta kriteria penerimaan. Desain menghasilkan diagram proses, relasi data, kontrak API, dan antarmuka. Implementasi menerjemahkan rancangan ke web Laravel dan aplikasi Flutter. Pengujian memeriksa fungsi, integrasi, dan penerimaan pengguna. Tahap pemeliharaan pada lingkup penelitian berupa perbaikan temuan, dokumentasi penggunaan, dan penyiapan prosedur pencadangan.

Urutan pekerjaan dipakai untuk menjaga hubungan antara kebutuhan, kode, dan pengujian. Apabila validasi pengguna mengubah kebutuhan, perubahan dicatat beserta alasannya, kemudian diagram, implementasi, dan skenario pengujian terkait diperbarui. Dengan demikian, dokumen akhir tetap merepresentasikan sistem yang diuji.''')
h('3.2.3 Analisis Kebutuhan Sistem',3)
p('Kebutuhan fungsional disusun menurut aktor dan proses. Kode kebutuhan digunakan untuk menelusuri kesesuaiannya terhadap skenario pengujian.')
func=[('F01','Publik','Melihat profil, kegiatan, berita, galeri, kebutuhan, dan kontak.'),('F02','Donatur biasa','Mengajukan donasi web tanpa akun dan membaca status melalui tautan rahasia.'),('F03','Donatur biasa','Mengajukan bantuan berisi beberapa barang melalui web.'),('F04','Donatur tetap','Mendaftar melalui mobile; akun menunggu persetujuan dan aktivasi email.'),('F05','Pengurus','Menyetujui, mengundang, menonaktifkan, dan memverifikasi kepemilikan riwayat akun.'),('F06','Donatur tetap','Login, memulihkan akses melalui email, mengubah profil dan kata sandi, serta logout.'),('F07','Donatur tetap','Membuat donasi mobile, membuka checkout, dan membaca status pembayaran.'),('F08','Donatur tetap','Melihat riwayat, detail, filter status, dan ringkasan donasi milik sendiri.'),('F09','Pengurus','Mengelola konten, kebutuhan, kategori, dan data donatur.'),('F10','Pengurus','Mencatat penerimaan donasi tunai atau transfer langsung.'),('F11','Pengurus','Mengonfirmasi, menolak, dan mencatat penerimaan bantuan barang.'),('F12','Pengurus','Mencatat, membukukan, membatalkan pengeluaran, serta mengatur izin bukti.'),('F13','Pengurus','Mengesahkan saldo pembukaan dan mengekspor laporan PDF atau Excel.'),('F14','Layanan transparansi','Menyajikan ringkasan dana, rincian pengeluaran dibukukan, dan bukti publik.'),('F15','Sistem','Memvalidasi webhook pembayaran dan mencegah pembaruan atau transaksi ganda.'),('F16','Sistem','Mencatat perubahan penting dan mengirim notifikasi sesuai konfigurasi.')]
table('Tabel 3.3 Kebutuhan fungsional',['Kode','Pengguna','Kebutuhan'],func,[1.4,3,9.6])
table('Tabel 3.4 Kebutuhan nonfungsional',['Aspek','Kriteria penerimaan yang direncanakan'],[('Keamanan akses','Akun tidak aktif tidak memperoleh akses pribadi; akun donatur tidak membuka panel pengurus; transaksi akun lain ditolak.'),('Konsistensi','Nilai pada web, API, mobile, PDF, dan Excel sama untuk periode serta status yang sama.'),('Integritas transaksi','Pengulangan permintaan dan webhook tidak menambah transaksi atau penerimaan yang sama dua kali.'),('Keandalan jaringan','Timeout tidak ditampilkan sebagai pembayaran berhasil; aplikasi menyediakan muat ulang dan mempertahankan referensi transaksi.'),('Kegunaan','Form mempunyai label dan pesan validasi; nominal serta status mudah dibaca; pengguna dapat menyelesaikan tugas inti saat UAT.'),('Privasi','Kontak pribadi dan bukti privat tidak muncul pada respons transparansi publik.'),('Pemeliharaan','Kode web dan mobile, kontrak API, migrasi, serta prosedur pencadangan dan pemulihan terdokumentasi.')],[3.2,10.8])
p('Lingkungan pengembangan meliputi komputer pengembang, server PHP dan MySQL, peramban, SDK Flutter, serta emulator atau perangkat Android. Versi dependensi mengikuti berkas proyek yang diuji. Pengujian integrasi membutuhkan jaringan yang menghubungkan perangkat ke backend dan alamat callback yang dapat dijangkau penyedia sandbox.')
h('3.3 Perancangan Sistem')
h('3.3.1 Use Case Diagram',3)
p('Aktor utama adalah pengunjung atau donatur biasa, donatur tetap, dan pengurus. Diagram dipisahkan menurut kelompok layanan agar hubungan fungsi tetap terbaca. Login menjadi prasyarat fungsi yang dilindungi. Midtrans, Fonnte, dan layanan email merupakan sistem eksternal pendukung yang dijelaskan pada arsitektur dan alur integrasi.')
fig('usecase-publik','Gambar 3.1 Use case layanan web publik',13)
fig('usecase-mobile','Gambar 3.2 Use case layanan mobile donatur tetap',13)
fig('usecase-pengurus','Gambar 3.3 Use case pengelolaan oleh pengurus',13)
p('Donatur tetap dapat melihat layanan transparansi melalui mobile setelah masuk ke aplikasi. Endpoint transparansi yang berisi informasi publik dapat diakses tanpa token; akses tersebut tetap tidak memberikan riwayat pribadi atau bukti privat. Pendaftaran akun belum memberikan kewenangan berdonasi sebagai pengguna terautentikasi sampai persetujuan dan aktivasi selesai.')
table('Tabel 3.5 Spesifikasi use case donasi mobile',['Elemen','Uraian'],[('Prasyarat','Akun donatur tetap aktif, email terverifikasi, dan sesi sah.'),('Masukan','Kategori donasi, nominal, dan kunci permintaan untuk mencegah duplikasi.'),('Alur utama','Memilih donasi, memvalidasi, menyimpan transaksi, membuka checkout, dan membaca status dari backend.'),('Alur alternatif','Validasi ditolak; jaringan gagal; pembayaran menunggu, gagal, kedaluwarsa, atau dibatalkan.'),('Pascakondisi','Transaksi tercatat dengan pemilik yang benar; penerimaan hanya diakui sesuai status pembayaran yang sah.')],[3.3,10.7])
h('3.3.2 Activity Diagram',3)
p('Diagram berikut memperlihatkan alur utama. Kondisi alternatif dijelaskan setelah masing-masing diagram agar keputusan saat terjadi kesalahan tetap menjadi bagian rancangan.')
fig('aktivitas-donasi','Gambar 3.4 Aktivitas donasi uang web dan mobile',12)
paras('''Apabila nominal atau kategori tidak valid, proses kembali ke formulir tanpa membuat checkout. Apabila permintaan ke penyedia pembayaran mengalami timeout, aplikasi mempertahankan referensi transaksi untuk pemeriksaan status. Redirect kembali dari halaman pembayaran tidak mengubah status menjadi Berhasil. Webhook yang tidak sah ditolak dan webhook berulang tidak menambah penerimaan dua kali. Pembatalan transaksi yang telah dikirim ke penyedia menunggu konfirmasi penyedia.

Riwayat web menggunakan tautan rahasia per transaksi. Riwayat mobile dibatasi pada pemilik yang sah. Donasi dari web tidak otomatis dimasukkan ke akun hanya berdasarkan kesamaan nomor WhatsApp; pengaitan transaksi memerlukan verifikasi pengurus.''')
fig('aktivitas-akun','Gambar 3.5 Aktivitas pendaftaran dan aktivasi akun',12)
p('Pengurus dapat menunda atau tidak menyetujui pengajuan yang belum memenuhi kriteria. Token aktivasi bersifat sekali pakai dan memiliki masa berlaku. Permintaan pemulihan akun menggunakan email terdaftar dengan respons yang tidak mengungkap keberadaan akun. Perubahan kata sandi atau pencabutan akses harus membatalkan sesi yang tidak lagi sah.')
fig('aktivitas-pengeluaran','Gambar 3.6 Aktivitas pembukuan pengeluaran',12)
paras('''Pengeluaran berstatus Draft dapat diperbaiki sebelum pembukuan dan belum memengaruhi saldo. Setelah dibukukan, koreksi dilakukan melalui mekanisme pembatalan yang mencatat pelaksana, waktu, dan alasan. Pengeluaran yang dibatalkan tidak lagi menjadi pengurang pada perhitungan laporan. Izin publikasi bukti ditetapkan secara terpisah dari status pembukuan.

Pada bantuan barang, donatur mengisi identitas dan daftar barang melalui web. Pengurus memeriksa pengajuan, kemudian mengonfirmasi atau menolaknya. Status Diterima diberikan setelah barang diterima secara fisik. Seluruh item dalam satu pengajuan dicatat diterima sekaligus; penyerahan parsial tidak termasuk rancangan.''')
h('3.3.3 Relasi Antar Tabel',3)
p('Basis data memisahkan profil donatur dari akun pengguna. Donatur biasa dapat memiliki profil tanpa akun, sedangkan donatur tetap dapat memiliki satu akun. Donasi menyimpan hubungan pemberi donasi serta pemilik akun yang berhak mengakses riwayat. Hubungan tersebut tidak menggantikan verifikasi kepemilikan transaksi.')
fig('erd-donasi','Gambar 3.7 Relasi data akun dan donasi',14)
fig('erd-bantuan','Gambar 3.8 Relasi data bantuan dan pengeluaran',14)
p('Notasi 1 : N menunjukkan satu record induk dapat terkait banyak record anak; 0..1 menunjukkan hubungan opsional dengan paling banyak satu record. Panah digunakan untuk menunjukkan arah pembacaan relasi, bukan aliran proses. Gambar memperlihatkan relasi inti, sedangkan atribut penelusuran dan data pendukung dirangkum pada tabel berikut.')
table('Tabel 3.6 Rancangan data inti dan pendukung',['Tabel','Data dan fungsi utama'],[('users','Identitas login, email, hash kata sandi, peran, status akun, dan verifikasi email.'),('donatur','Nama, kontak, jenis donatur, dan user_id opsional yang unik.'),('donasi','Kode transaksi, donatur, kategori, pemilik, nominal, biaya, sumber, metode, status, serta waktu pembayaran.'),('payment_transactions','Hubungan unik ke donasi, referensi penyedia, checkout, dan informasi status pembayaran.'),('bantuan dan bantuan_items','Pengajuan bantuan, donatur, status, serta daftar nama barang, jumlah, satuan, dan kategori.'),('pengeluaran','Tanggal, kategori, keperluan, nominal, penerima, status, bukti, izin publikasi, pencatat, pembuku, dan pembatal.'),('saldo_pembukaan','Tanggal mulai, nominal, keterangan, dan pengurus yang mengesahkan.'),('donor_account_audits','Jejak tindakan pengurus atau pengguna terkait akun dan kepemilikan transaksi.'),('status_histories','Riwayat perubahan status transaksi untuk penelusuran proses.'),('kebutuhan dan konten','Target serta realisasi kebutuhan; profil, kegiatan, berita, galeri, dan pengaturan publikasi.'),('personal_access_tokens','Token akses API; tidak ditampilkan pada laporan atau antarmuka publik.')],[4.3,9.7])
p('Nominal uang menggunakan tipe desimal agar pembulatan dapat dikendalikan. Kode transaksi dan hubungan pembayaran diberi batasan unik. Foreign key digunakan untuk menjaga hubungan data. Data transaksi tidak dihapus hanya untuk memperbaiki status; perubahan yang memengaruhi laporan harus dapat ditelusuri.')
h('3.3.4 Arsitektur Sistem',3)
fig('arsitektur','Gambar 3.9 Arsitektur web dan mobile',14)
paras('''Web publik dan panel pengurus mengirim permintaan ke backend Laravel. Aplikasi Flutter menggunakan endpoint /api/v1 untuk autentikasi, referensi, donasi, status pembayaran, profil, dan transparansi. Backend menerapkan validasi serta kewenangan sebelum mengakses MySQL. Perangkat mobile tidak mengakses basis data secara langsung.

Integrasi Midtrans dilakukan melalui backend; rahasia pembayaran tidak disimpan pada aplikasi mobile. Checkout ditampilkan melalui halaman pembayaran, sedangkan status transaksi ditentukan dari informasi penyedia yang diverifikasi backend. Fonnte digunakan untuk notifikasi WhatsApp sesuai konfigurasi. Layanan email digunakan untuk aktivasi dan pemulihan akun. Kegagalan pengiriman notifikasi tidak membatalkan transaksi yang telah sah.''')
table('Tabel 3.7 Kelompok antarmuka API',['Kelompok','Contoh endpoint','Pengendalian'],[('Autentikasi','POST /auth/register; POST /auth/login','Validasi, pembatasan percobaan, dan status akun.'),('Profil','GET /profile; PUT /profile','Token sah dan identitas pengguna sendiri.'),('Donasi','GET/POST /donations; GET /donations/{id}','Kepemilikan pada backend dan pencegahan pengajuan ganda.'),('Pembayaran','POST /donations/{id}/pay; GET /donations/{id}/status','Pemilik transaksi dan hasil verifikasi server.'),('Transparansi','GET /transparency/summary; GET /transparency/expenses','Data publik terpilih; bukti privat ditolak.')],[2.8,6.2,5])
h('3.3.4.1 Aturan Perhitungan Dana',3)
paras('''Penerimaan bruto adalah jumlah nominal donasi berstatus Berhasil. Penerimaan bersih diperoleh dari donasi manual yang dihitung penuh dan donasi online berhasil setelah dikurangi biaya yang telah tercatat. Donasi online dengan biaya belum tersedia tetap dapat ditampilkan pada bruto, tetapi belum dimasukkan ke penerimaan bersih. Sistem memberi penanda bahwa saldo masih sementara, termasuk bila biaya yang belum tersedia berasal dari transaksi sebelum periode laporan.

Saldo awal periode berasal dari saldo pembukaan yang berlaku, ditambah penerimaan bersih sebelum awal periode dan dikurangi pengeluaran dibukukan sebelum awal periode. Saldo akhir periode diperoleh dengan menambahkan penerimaan bersih selama periode ke saldo awal, kemudian mengurangi pengeluaran dibukukan selama periode tersebut. Penerimaan memakai tanggal pembayaran, sedangkan pengeluaran memakai tanggal transaksi pengeluaran.

Sebagai contoh pengujian sintetis, saldo awal Rp1.000.000, donasi manual Rp100.000, donasi online berhasil Rp200.000 dengan biaya Rp2.000, dan pengeluaran dibukukan Rp50.000 menghasilkan penerimaan bersih Rp298.000 serta saldo akhir Rp1.248.000. Donasi online lain sebesar Rp150.000 dengan biaya belum tersedia menambah bruto, tetapi belum menambah penerimaan bersih. Angka ini hanya contoh uji, bukan data keuangan panti.

Biaya pada sandbox diberi keterangan simulasi dan tidak dianggap biaya produksi. Saldo yang disajikan merupakan saldo berdasarkan transaksi tercatat, sehingga tidak dinyatakan sebagai saldo rekening bank yang telah direkonsiliasi.''')
h('3.3.5 Antarmuka Pengguna',3)
p('Rancangan antarmuka berikut menampilkan susunan fungsi, bukan data transaksi nyata. Istilah dan status dibuat konsisten antara web dan mobile. Pesan kesalahan ditempatkan dekat masukan terkait, sedangkan tindakan pembukuan dan pembatalan memerlukan konfirmasi yang menjelaskan akibatnya.')
fig('ui-web','Gambar 3.10 Rancangan halaman web publik',14)
p('Halaman publik menempatkan informasi panti, kebutuhan, dan tindakan donasi pada navigasi utama. Form donasi meminta identitas yang diperlukan, kategori, dan nominal. Setelah pengajuan, pengguna memperoleh akses status melalui tautan transaksi. Form bantuan memungkinkan beberapa barang dengan nama, jumlah, dan satuan yang jelas.')
fig('ui-admin','Gambar 3.11 Rancangan panel pengurus',14)
p('Panel pengurus menyajikan ringkasan dan daftar transaksi dengan filter periode serta status. Halaman pengeluaran memisahkan tindakan menyimpan Draft, membukukan, dan membatalkan. Halaman laporan menyajikan komponen bruto, biaya, penerimaan bersih, pengeluaran, serta saldo agar pengurus dapat menelusuri perhitungannya sebelum ekspor.')
fig('ui-mobile','Gambar 3.12 Rancangan antarmuka aplikasi Android',14)
p('Mobile menggunakan navigasi beranda, riwayat, transparansi, dan profil. Layar pendaftaran menjelaskan status menunggu persetujuan. Detail donasi menampilkan nominal, kategori, referensi, dan status dari server. Transparansi menyediakan filter periode dan rincian pengeluaran; tombol bukti hanya tersedia untuk bukti yang diizinkan. Saat jaringan terputus, aplikasi menampilkan pesan dan tindakan muat ulang.')
h('3.4 Pengujian Sistem')
h('3.4.1 Pengujian Fungsional dan Hak Akses',3)
p('Pengujian dilakukan menggunakan data uji yang tidak bercampur dengan catatan operasional. Setiap kasus dicatat dengan masukan, hasil yang diharapkan, hasil aktual, bukti, dan status lulus atau gagal. Tabel berikut memuat rencana kasus utama; belum menyatakan hasil pelaksanaan.')
tests=[('U01 / F01,F09','Publikasikan dan tarik konten','Konten hanya tampil pada publik jika berstatus terbit.'),('U02 / F02','Donasi web valid dan nominal tidak valid','Transaksi valid tersimpan; masukan salah ditolak.'),('U03 / F02,F08','Gunakan tautan atau ID transaksi orang lain','Akses tanpa token rahasia atau kepemilikan sah ditolak.'),('U04 / F04,F05','Daftar, setujui, dan tukarkan aktivasi','Akun menunggu tidak dapat login; aktif setelah token sah.'),('U05 / F06','Reset, ubah kata sandi, lalu pakai sesi lama','Pemulihan sah berhasil; sesi yang dicabut tidak berlaku.'),('U06 / F07,F15','Kirim ulang pembuatan donasi yang sama','Tidak terbentuk transaksi atau checkout ganda.'),('U07 / F07,F15','Webhook sah, palsu, nominal salah, dan berulang','Hanya pemberitahuan sah memperbarui status satu kali.'),('U08 / F07','Tutup checkout atau kirim redirect sukses palsu','Status mengikuti backend, bukan redirect klien.'),('U09 / F07','Pembatalan pembayaran dengan timeout provider','Tidak ditandai batal sebelum kepastian yang sah.'),('U10 / F03,F11','Ajukan banyak barang lalu catat penerimaan','Item lengkap dan status sesuai konfirmasi fisik.'),('U11 / F10,F12','Donasi manual, Draft, pembukuan berulang, pembatalan','Donasi manual dihitung penuh; pengeluaran memengaruhi saldo sesuai status.'),('U12 / F13,F14','Bandingkan web, mobile, PDF, dan Excel','Komponen dana sama pada periode yang sama.'),('U13 / F13,F14','Biaya kosong pada periode kini dan sebelumnya','Bruto tetap tampil; saldo bersih diberi penanda sementara.'),('U14 / F14','Buka bukti publik dan bukti privat','Hanya bukti yang diizinkan dapat dibuka.'),('U15 / F05,F06','Akun donatur mencoba panel pengurus','Panel ditolak; hak peran tidak dapat diubah dari klien.'),('U16 / F16','Provider notifikasi gagal atau timeout','Hasil pengiriman tercatat tanpa menggandakan transaksi.')]
table('Tabel 3.8 Rencana pengujian fungsi',['Kode / kebutuhan','Skenario','Hasil yang diharapkan'],tests,[2.4,5.7,5.9])
h('3.4.2 Pengujian Integrasi',3)
paras('''Pengujian integrasi dilakukan pada rangkaian pendaftaran–persetujuan–aktivasi–login, donasi–checkout–webhook–riwayat, serta pengeluaran–pembukuan–laporan–transparansi. Pemeriksaan membandingkan data sumber, respons API, tampilan aplikasi, dan hasil ekspor. Transaksi sekitar pergantian bulan diuji untuk memastikan batas periode konsisten.

Pengujian backend dengan respons penyedia yang disimulasikan dipisahkan dari pengujian langsung Midtrans sandbox. Pengujian langsung harus mencatat referensi sandbox, waktu, status penyedia, dan status aplikasi. Keberhasilan tes unit atau widget belum menjadi bukti keberhasilan seluruh alur web dan mobile. Perangkat Android dan kondisi jaringan yang digunakan dicatat dalam laporan pengujian.''')
h('3.4.3 Pengujian Penerimaan Pengguna',3)
paras('''Pengujian penerimaan melibatkan pengurus dan perwakilan donatur yang bersedia. Pemilihan peserta dilakukan berdasarkan keterlibatan dalam administrasi atau penggunaan layanan donasi. Jumlah peserta aktual, peran, dan pengalaman penggunaan dicatat saat pelaksanaan; proposal ini tidak menetapkan hasil atau jumlah responden yang belum diperoleh.

Pengurus diminta mengelola konten, mencatat donasi manual, menerima bantuan, membukukan pengeluaran, dan menghasilkan laporan. Donatur biasa mencoba form web serta status transaksi. Donatur tetap mencoba login, donasi mobile, riwayat, dan transparansi. Peneliti mencatat keberhasilan tugas, kesalahan, bantuan yang diperlukan, serta komentar peserta.

Tanggapan dikumpulkan melalui pertanyaan mengenai kejelasan menu, kemudahan pengisian, keterbacaan status, pemahaman ringkasan dana, dan kesesuaian fungsi. Skala 1 sampai 5 dapat digunakan untuk merangkum tanggapan secara deskriptif, disertai jumlah responden dan komentar. Skor tersebut tidak diberi label SUS atau TAM karena instrumennya berbeda. Semua kegagalan kritis pada hak akses, pembayaran, dan perhitungan saldo harus diperbaiki serta diuji ulang sebelum sistem dinyatakan memenuhi kriteria penerimaan.''')

chapter('BAB IV\nJADWAL KEGIATAN')
h('4.1 Jadwal')
paras('''Penelitian direncanakan selama enam bulan. Bulan pertama dihitung sejak jadwal pelaksanaan disetujui pembimbing. Penetapan bulan kalender mengikuti agenda akademik dan kesediaan pihak panti. Penyusunan laporan berjalan bersama pengembangan agar perubahan kebutuhan dan hasil pengujian tercatat.''')
schedule=[('Identifikasi masalah dan studi pustaka',[1]),('Observasi, wawancara, dokumentasi',[1,2]),('Analisis dan validasi kebutuhan',[2]),('Perancangan proses, data, API, antarmuka',[2,3]),('Pengembangan web dan akun donatur',[3,4]),('Donasi, bantuan, pengeluaran, laporan',[3,4]),('Pengembangan Flutter dan integrasi API',[4,5]),('Pengujian fungsi dan integrasi',[5]),('UAT dan perbaikan',[5,6]),('Penyusunan laporan',[2,3,4,5,6]),('Finalisasi dan demonstrasi',[6])]
t=table('Tabel 4.1 Gantt chart rencana penelitian',['Kegiatan','B1','B2','B3','B4','B5','B6'],[(a,*['X' if m in ms else '' for m in range(1,7)]) for a,ms in schedule],[7,1.15,1.15,1.15,1.15,1.15,1.15])
for ri,(_,ms) in enumerate(schedule,1):
 for m in ms:
  pr=t.rows[ri].cells[m]._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'D9D9D9');pr.append(sh)
p('Keterangan: X dan arsiran menunjukkan bulan pelaksanaan. B1 sampai B6 merupakan bulan relatif, bukan klaim bahwa kegiatan sudah selesai.')
h('4.2 Pembagian Tugas')
paras('''Penelitian dilaksanakan secara individu oleh Noor Achmad Rozaq Ramdhani. Pengurus berperan sebagai narasumber dan pengguna penguji, sedangkan pembimbing memberikan arahan akademik. Keterlibatan pihak panti dan peserta UAT dijadwalkan sesuai kesediaan mereka.''')
table('Tabel 4.2 Tanggung jawab dan keluaran',['Pelaksana','Tanggung jawab','Keluaran'],[('Peneliti','Pengumpulan data, analisis, desain, implementasi web dan mobile, pengujian, serta dokumentasi.','Proposal, rancangan, kode aplikasi, bukti uji, panduan penggunaan, dan laporan akhir.'),('Pengurus panti','Mengonfirmasi kebutuhan, aturan penerimaan dan pengeluaran, batas publikasi, serta mencoba sistem.','Masukan kebutuhan dan hasil evaluasi tugas pengurus.'),('Perwakilan donatur','Mencoba layanan web atau mobile sesuai peran dan memberi tanggapan.','Catatan keberhasilan tugas, kendala, dan saran penggunaan.'),('Dosen pembimbing','Meninjau ruang lingkup, metode, rancangan, serta penyajian hasil penelitian.','Arahan dan catatan revisi akademik.')],[3,5.5,5.5])
p('Keluaran akhir yang ditargetkan adalah aplikasi web dan Android yang terhubung, rancangan data serta API, laporan pengujian dengan bukti, dan panduan penggunaan. Penyerahan untuk penggunaan operasional dilakukan setelah konfigurasi layanan, data resmi, prosedur pencadangan, dan hasil penerimaan pengguna ditinjau bersama pihak panti.')

chapter('DAFTAR PUSTAKA')
refs=[
'Chandra, Hendrik, J., & Leonardo. (2022). Perancangan sistem informasi geografis dan donasi online panti asuhan di Kota Medan menggunakan metode WebQual 4.0. Jurnal TIMES, 11(1), 1–8.',
'Destriani, M., & Bahri, M. I. S. (2026). Sistem informasi pengelolaan donasi pada Panti Asuhan Al-Rasyid. Global, 13(1), 14–30.',
'Kusnadi, M. P., & Daniawan, B. (2024). Sistem informasi manajemen donasi pada panti asuhan menggunakan metode pengujian Technology Acceptance Model (TAM). Jurnal INSTEK (Informatika Sains dan Teknologi), 9(2), 336–345.',
'Mauliana, P., Firmansyah, R., & Sutardi, A. (2019). Perancangan sistem informasi panti asuhan “e-Panti” berbasis web. JUST IT: Jurnal Sistem Informasi, Teknologi Informatika dan Komputer, 10(1), 62–71. https://doi.org/10.24853/justit.10.1.62-71',
'Romansyah, J., Mulyani, A., & Kurniadi, D. (2019). Rancang bangun sistem informasi penggalangan dana sosial panti asuhan berbasis web. Jurnal Algoritma, 16(2), 158–165.',
'Safira, A., Suryantari, P. A., & Ramadhan, J. B. (2024). Perancangan sistem informasi donasi panti asuhan berbasis web Panti Peduli. Prosiding Seminar Nasional Teknologi dan Sistem Informasi (SITASI), 351–364.',
'Salim, R. R. M. (2020). Rancang bangun sistem informasi administrasi untuk panti asuhan. Jurnal Ilmiah MATRIK, 22(2), 195–200.',
'Sintiya, E. S., Vista, C. B., Wibowo, D. W., Kusumawardana, A., & Puspitaningrum, A. C. (2025). Rancang bangun sistem informasi berbasis extreme programming untuk transparansi pengelolaan donasi panti asuhan. JATI (Jurnal Mahasiswa Teknik Informatika), 9(1), 291–297.',
'Ujianti, M. H., & Al-Farisy, M. F. (2026). Analisis sistem pendataan anak asuh di Panti Asuhan Muhammadiyah Slawi. ZONAsi: Jurnal Sistem Informasi, 8(1), 275–281.',
'Yunita, L., Neneng, Isnain, A. R., & Dellia, P. (2022). Analisis perancangan sistem informasi akuntansi pencatatan dan pengelolaan keuangan pada Yayasan Panti Asuhan Harapan Karomah. Jurnal Ilmiah Sistem Informasi Akuntansi, 2(2), 62–68.'
]
for t in refs:
 q=p();q.alignment=WD_ALIGN_PARAGRAPH.LEFT;q.paragraph_format.left_indent=Cm(1);q.paragraph_format.first_line_indent=Cm(-1);q.paragraph_format.line_spacing=1;q.paragraph_format.space_after=Pt(12)
 # APA: periodical title and volume in italics, issue and pages roman.
 import re
 m=re.search(r'\. (Jurnal TIMES|Global|Jurnal INSTEK|JUST IT:|Jurnal Algoritma|Prosiding Seminar|Jurnal Ilmiah MATRIK|JATI |ZONAsi:|Jurnal Ilmiah Sistem Informasi Akuntansi)',t)
 if m:
  a=m.start()+2;b=t.find('(',a)
  if t[a:].startswith(('Jurnal INSTEK','JATI ')):b=t.find('(',t.find('),',a)+2)
  if t[a:].startswith('Prosiding'):b=t.find(', 351',a)
  if b<0:b=len(t)
  q.add_run(t[:a]);q.add_run(t[a:b]).italic=True;q.add_run(t[b:])
 else:q.add_run(t)
doc.core_properties.title=TITLE;doc.core_properties.author='Noor Achmad Rozaq Ramdhani';doc.core_properties.subject='Proposal Laporan Akhir D3 Manajemen Informatika'
# Remove theme-font and title-border defaults from the packaged Word template.
for st in doc.styles:
 if st.type in (1,2):
  st.font.name='Times New Roman'
  for rf in st.element.iter(qn('w:rFonts')):
   for attr in list(rf.attrib):
    if 'Theme' in attr:del rf.attrib[attr]
  for borders in list(st.element.iter(qn('w:pBdr'))):borders.getparent().remove(borders)
for q in doc.paragraphs:
 for r in q.runs:
  r.font.name='Times New Roman'
 for borders in list(q._p.iter(qn('w:pBdr'))):borders.getparent().remove(borders)
doc.save(OUT/'Proposal_LA_Donasi_Web_Mobile_An_Nuur_2.docx')
print(OUT/'Proposal_LA_Donasi_Web_Mobile_An_Nuur_2.docx')
